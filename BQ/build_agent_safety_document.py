from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


OUTPUT = Path("Agent_Safety_Architecture_Guardrails_and_Gatekeepers.docx").resolve()
PLANNING_DIAGRAM_PATH = Path("/private/tmp/bq-agent-safety-planning-flow.png")
DISPENSE_DIAGRAM_PATH = Path("/private/tmp/bq-agent-safety-dispense-flow.png")

NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "5B6573"
PALE_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
PALE_GOLD = "FFF4CE"
GOLD = "7A5A00"
PALE_RED = "FDEBEC"
RED = "9B1C1C"
PALE_GREEN = "EAF4EA"
GREEN = "245B2A"
WHITE = "FFFFFF"
BLACK = "1A1A1A"


def _diagram_font(size: int, bold: bool = False):
    path = "/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Arial.ttf"
    return ImageFont.truetype(path, size=size)


def _draw_centered_text(draw, box, text, font, fill, spacing=5):
    x1, y1, x2, y2 = box
    lines = text.split("\n")
    heights = []
    widths = []
    for line in lines:
        bounds = draw.textbbox((0, 0), line, font=font)
        widths.append(bounds[2] - bounds[0])
        heights.append(bounds[3] - bounds[1])
    total_h = sum(heights) + spacing * (len(lines) - 1)
    y = y1 + ((y2 - y1) - total_h) / 2
    for line, width, height in zip(lines, widths, heights):
        draw.text((x1 + ((x2 - x1) - width) / 2, y), line, font=font, fill=fill)
        y += height + spacing


def _draw_box(draw, box, text, *, fill, outline, font, text_fill=NAVY, radius=18, width=3):
    draw.rounded_rectangle(box, radius=radius, fill="#" + fill, outline="#" + outline, width=width)
    _draw_centered_text(draw, box, text, font, "#" + text_fill)


def _draw_labeled_box(
    draw,
    box,
    title,
    body,
    *,
    fill,
    outline,
    title_font,
    body_font,
    text_fill=NAVY,
    radius=22,
    width=4,
):
    draw.rounded_rectangle(box, radius=radius, fill="#" + fill, outline="#" + outline, width=width)
    x1, y1, x2, y2 = box
    title_bounds = draw.textbbox((0, 0), title, font=title_font)
    title_height = title_bounds[3] - title_bounds[1]
    body_lines = body.split("\n") if body else []
    body_heights = []
    for line in body_lines:
        bounds = draw.textbbox((0, 0), line, font=body_font)
        body_heights.append(bounds[3] - bounds[1])
    body_height = sum(body_heights) + max(0, len(body_lines) - 1) * 7
    gap = 15 if body_lines else 0
    total_height = title_height + gap + body_height
    y = y1 + ((y2 - y1) - total_height) / 2
    title_width = title_bounds[2] - title_bounds[0]
    draw.text((x1 + ((x2 - x1) - title_width) / 2, y), title, font=title_font, fill="#" + text_fill)
    y += title_height + gap
    for line, line_height in zip(body_lines, body_heights):
        bounds = draw.textbbox((0, 0), line, font=body_font)
        line_width = bounds[2] - bounds[0]
        draw.text((x1 + ((x2 - x1) - line_width) / 2, y), line, font=body_font, fill="#" + text_fill)
        y += line_height + 7


def _draw_arrow(draw, start, end, *, color=BLUE, width=6, head=16, label=None, label_font=None):
    draw.line([start, end], fill="#" + color, width=width)
    x1, y1 = start
    x2, y2 = end
    if abs(x2 - x1) >= abs(y2 - y1):
        direction = 1 if x2 >= x1 else -1
        points = [(x2, y2), (x2 - direction * head, y2 - head * 0.65), (x2 - direction * head, y2 + head * 0.65)]
    else:
        direction = 1 if y2 >= y1 else -1
        points = [(x2, y2), (x2 - head * 0.65, y2 - direction * head), (x2 + head * 0.65, y2 - direction * head)]
    draw.polygon(points, fill="#" + color)
    if label and label_font:
        bounds = draw.textbbox((0, 0), label, font=label_font)
        lx = (x1 + x2) / 2 - (bounds[2] - bounds[0]) / 2
        ly = (y1 + y2) / 2 - (bounds[3] - bounds[1]) - 6
        draw.rounded_rectangle((lx - 8, ly - 3, lx + (bounds[2] - bounds[0]) + 8, ly + (bounds[3] - bounds[1]) + 7), radius=7, fill="#FFFFFF")
        draw.text((lx, ly), label, font=label_font, fill="#" + color)


def create_safety_diagrams(planning_path: Path, dispense_path: Path) -> tuple[Path, Path]:
    title_font = _diagram_font(38, bold=True)
    box_title_font = _diagram_font(27, bold=True)
    box_body_font = _diagram_font(23, bold=False)
    decision_font = _diagram_font(24, bold=True)
    label_font = _diagram_font(21, bold=True)

    # Figure 1: protocol and order approval.
    planning = Image.new("RGB", (1800, 620), "#FFFFFF")
    draw = ImageDraw.Draw(planning)
    draw.rounded_rectangle((20, 18, 1780, 600), radius=28, fill="#FFFFFF", outline="#D7DDE5", width=3)
    draw.text((55, 42), "PROTOCOL AND ORDER AUTHORIZATION", font=title_font, fill="#" + NAVY)
    draw.line((55, 102, 1745, 102), fill="#D7DDE5", width=3)

    proposal = (45, 175, 250, 355)
    guardrail = (310, 175, 615, 355)
    protocol = (675, 160, 1080, 370)
    diamond = [(1230, 160), (1360, 265), (1230, 370), (1100, 265)]
    order = (1430, 160, 1755, 370)
    negative = (935, 450, 1525, 565)

    _draw_labeled_box(draw, proposal, "AGENT", "Proposed\nprotocol", fill=LIGHT_GRAY, outline="B7C0CC", title_font=box_title_font, body_font=box_body_font)
    _draw_labeled_box(draw, guardrail, "GUARDRAIL", "Structure proposal\nCollect missing data", fill=PALE_BLUE, outline=BLUE, title_font=box_title_font, body_font=box_body_font)
    _draw_labeled_box(draw, protocol, "PROTOCOL SAFETY GATE", "Conditions  •  Medications\nDose  •  Timing", fill=PALE_BLUE, outline=BLUE, title_font=box_title_font, body_font=box_body_font)
    draw.polygon(diamond, fill="#FFF4CE", outline="#" + GOLD)
    draw.line(diamond + [diamond[0]], fill="#" + GOLD, width=4, joint="curve")
    _draw_centered_text(draw, (1100, 160, 1360, 370), "ALL CHECKS\nPASS?", decision_font, "#" + GOLD, spacing=6)
    _draw_labeled_box(draw, order, "ORDER GATEKEEPER", "Signed authorization\nCartridge manifest", fill=PALE_GREEN, outline=GREEN, title_font=box_title_font, body_font=box_body_font, text_fill=GREEN)
    _draw_labeled_box(draw, negative, "NO APPROVAL", "Request input  •  Replan  •  Clinical review  •  Reject", fill=PALE_RED, outline=RED, title_font=box_title_font, body_font=box_body_font, text_fill=RED)

    _draw_arrow(draw, (250, 265), (310, 265), color=BLUE)
    _draw_arrow(draw, (615, 265), (675, 265), color=BLUE)
    _draw_arrow(draw, (1080, 265), (1100, 265), color=BLUE)
    _draw_arrow(draw, (1360, 265), (1430, 265), color=GREEN, label="YES", label_font=label_font)
    _draw_arrow(draw, (1230, 370), (1230, 450), color=RED, label="NO", label_font=label_font)
    planning.save(planning_path, format="PNG", optimize=True, dpi=(220, 220))

    # Figure 2: independent checks immediately before physical release.
    dispense = Image.new("RGB", (1800, 650), "#FFFFFF")
    draw = ImageDraw.Draw(dispense)
    draw.rounded_rectangle((20, 18, 1780, 630), radius=28, fill="#FFFFFF", outline="#D7DDE5", width=3)
    draw.text((55, 42), "DISPENSE-TIME DEFENSE", font=title_font, fill="#" + NAVY)
    draw.line((55, 102, 1745, 102), fill="#D7DDE5", width=3)

    inputs = (45, 220, 320, 430)
    cartridge = (405, 135, 900, 295)
    pill = (405, 355, 900, 515)
    diamond = [(1070, 220), (1210, 325), (1070, 430), (930, 325)]
    release = (1300, 220, 1755, 430)
    blocked = (1280, 505, 1690, 605)

    _draw_labeled_box(draw, inputs, "USER + CARTRIDGE", "Authenticated user\nInstalled cartridge", fill=LIGHT_GRAY, outline="B7C0CC", title_font=box_title_font, body_font=box_body_font)
    _draw_labeled_box(draw, cartridge, "CARTRIDGE GATEKEEPER", "User match  •  Manifest  •  DHR SAFE\nLot status  •  Expiry", fill=PALE_BLUE, outline=BLUE, title_font=box_title_font, body_font=box_body_font)
    _draw_labeled_box(draw, pill, "PER-PILL GATEKEEPER", "Current profile  •  Interactions\nDaily dose  •  Actual timing", fill=PALE_BLUE, outline=BLUE, title_font=box_title_font, body_font=box_body_font)
    draw.polygon(diamond, fill="#FFF4CE", outline="#" + GOLD)
    draw.line(diamond + [diamond[0]], fill="#" + GOLD, width=4, joint="curve")
    _draw_centered_text(draw, (930, 220, 1210, 430), "BOTH GATES\nALLOW?", decision_font, "#" + GOLD, spacing=6)
    _draw_labeled_box(draw, release, "DISPENSE ONE PILL", "Single-use authorization\nconsumed atomically", fill=PALE_GREEN, outline=GREEN, title_font=box_title_font, body_font=box_body_font, text_fill=GREEN)
    _draw_labeled_box(draw, blocked, "HARD BLOCK", "No physical release", fill=PALE_RED, outline=RED, title_font=box_title_font, body_font=box_body_font, text_fill=RED)

    draw.line((320, 325, 360, 325), fill="#" + BLUE, width=6)
    draw.line((360, 215, 360, 435), fill="#" + BLUE, width=6)
    _draw_arrow(draw, (360, 215), (405, 215), color=BLUE)
    _draw_arrow(draw, (360, 435), (405, 435), color=BLUE)
    _draw_arrow(draw, (900, 215), (965, 270), color=BLUE)
    _draw_arrow(draw, (900, 435), (965, 380), color=BLUE)
    _draw_arrow(draw, (1210, 325), (1300, 325), color=GREEN, label="YES", label_font=label_font)
    draw.line((1070, 430, 1070, 555), fill="#" + RED, width=6)
    _draw_arrow(draw, (1070, 555), (1280, 555), color=RED, label="NO", label_font=label_font)
    dispense.save(dispense_path, format="PNG", optimize=True, dpi=(220, 220))

    return planning_path, dispense_path


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = tc_mar.find(qn(f"w:{side}"))
        if tag is None:
            tag = OxmlElement(f"w:{side}")
            tc_mar.append(tag)
        tag.set(qn("w:w"), str(value))
        tag.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_table_geometry(table, widths_dxa: list[int], indent_dxa: int = 120) -> None:
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[min(idx, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_borders(table, color="CBD2D9", size="6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_run_font(run, size=None, color=BLACK, bold=None, italic=None, name="Calibri") -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_keep_with_next(paragraph, value=True) -> None:
    paragraph.paragraph_format.keep_with_next = value


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = instruction
    fld_char_separate = OxmlElement("w:fldChar")
    fld_char_separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_begin, instr_text, fld_char_separate, text, fld_char_end])
    set_run_font(run, size=9, color=MUTED)


def add_custom_numbering(doc: Document) -> tuple[int, int, int]:
    numbering = doc.part.numbering_part.element
    existing_abs = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    existing_num = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    next_abs = max(existing_abs, default=0) + 1
    next_num = max(existing_num, default=0) + 1

    def make_abstract(abs_id: int, ordered: bool) -> None:
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abs_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), "decimal" if ordered else "bullet")
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), "%1." if ordered else "•")
        lvl_jc = OxmlElement("w:lvlJc")
        lvl_jc.set(qn("w:val"), "left")
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "720")
        tabs.append(tab)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "720")
        ind.set(qn("w:hanging"), "360")
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "160")
        spacing.set(qn("w:line"), "280")
        spacing.set(qn("w:lineRule"), "auto")
        p_pr.extend([tabs, ind, spacing])
        lvl.extend([start, num_fmt, lvl_text, lvl_jc, p_pr])
        if not ordered:
            r_pr = OxmlElement("w:rPr")
            fonts = OxmlElement("w:rFonts")
            fonts.set(qn("w:ascii"), "Symbol")
            fonts.set(qn("w:hAnsi"), "Symbol")
            r_pr.append(fonts)
            lvl.append(r_pr)
        abstract.append(lvl)
        numbering.append(abstract)

    make_abstract(next_abs, False)
    bullet_num = next_num
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(bullet_num))
    abs_ref = OxmlElement("w:abstractNumId")
    abs_ref.set(qn("w:val"), str(next_abs))
    num.append(abs_ref)
    numbering.append(num)

    make_abstract(next_abs + 1, True)
    decimal_num = next_num + 1
    num2 = OxmlElement("w:num")
    num2.set(qn("w:numId"), str(decimal_num))
    abs_ref2 = OxmlElement("w:abstractNumId")
    abs_ref2.set(qn("w:val"), str(next_abs + 1))
    num2.append(abs_ref2)
    numbering.append(num2)
    decimal_num_restart = next_num + 2
    num3 = OxmlElement("w:num")
    num3.set(qn("w:numId"), str(decimal_num_restart))
    abs_ref3 = OxmlElement("w:abstractNumId")
    abs_ref3.set(qn("w:val"), str(next_abs + 1))
    num3.append(abs_ref3)
    level_override = OxmlElement("w:lvlOverride")
    level_override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    level_override.append(start_override)
    num3.append(level_override)
    numbering.append(num3)
    return bullet_num, decimal_num, decimal_num_restart


def add_list_item(doc, text: str, num_id: int, *, bold_prefix: str | None = None):
    p = doc.add_paragraph()
    p_pr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_el])
    p_pr.append(num_pr)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, size=11, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, size=11)
    else:
        r = p.add_run(text)
        set_run_font(r, size=11)
    return p


def add_callout(doc, label: str, text: str, *, fill=PALE_BLUE, accent=BLUE, color=NAVY):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.line_spacing = 1.1
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), accent)
    borders.append(left)
    p_pr.append(borders)
    r1 = p.add_run(label.upper() + "  ")
    set_run_font(r1, size=10, color=accent, bold=True)
    r2 = p.add_run(text)
    set_run_font(r2, size=10.5, color=color)
    return p


def add_para(doc, text: str, *, bold=False, italic=False, color=BLACK, after=6, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.1
    p.paragraph_format.widow_control = True
    p.paragraph_format.keep_with_next = keep
    r = p.add_run(text)
    set_run_font(r, size=11, color=color, bold=bold, italic=italic)
    return p


def add_table(doc, headers: list[str], rows: list[list[str]], widths: list[int], *, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, LIGHT_GRAY)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(header)
        set_run_font(r, size=9.5, color=NAVY, bold=True)
    for row_idx, values in enumerate(rows):
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        for idx, value in enumerate(values):
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(value)
            set_run_font(r, size=font_size, color=BLACK)
            if row_idx % 2 == 1:
                set_cell_shading(cells[idx], "FAFBFC")
    set_table_geometry(table, widths)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)
    spacer.paragraph_format.space_before = Pt(0)
    return table


def add_page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    normal.paragraph_format.widow_control = True

    heading_tokens = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    for name, size, color, bold in (
        ("Document Title", 25, NAVY, True),
        ("Document Subtitle", 13.5, MUTED, False),
        ("Kicker", 9.5, BLUE, True),
    ):
        if name not in styles:
            style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        else:
            style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = bold
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(5 if name == "Document Title" else 12)
        style.paragraph_format.keep_with_next = True


def configure_page(section) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.78)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.42)
    section.footer_distance = Inches(0.42)


def configure_header_footer(section) -> None:
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    p_pr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "D7DDE5")
    borders.append(bottom)
    p_pr.append(borders)
    r = p.add_run("AGENT SAFETY ARCHITECTURE  |  GUARDRAILS & GATEKEEPERS")
    set_run_font(r, size=8.5, color=MUTED, bold=True)

    footer = section.footer
    table = footer.add_table(rows=1, cols=2, width=Inches(6.5))
    set_table_geometry(table, [7200, 2160], indent_dxa=0)
    # Remove all borders from footer table.
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "nil")
        borders.append(tag)
    tbl_pr.append(borders)
    left = table.cell(0, 0).paragraphs[0]
    left.paragraph_format.space_after = Pt(0)
    lr = left.add_run("System specification  |  Internal working document")
    set_run_font(lr, size=8.5, color=MUTED)
    right = table.cell(0, 1).paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    right.paragraph_format.space_after = Pt(0)
    rr = right.add_run("Page ")
    set_run_font(rr, size=8.5, color=MUTED)
    add_field(right, "PAGE")


def heading(doc, text: str, level: int = 1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.widow_control = True
    return p


def build_document() -> Path:
    planning_diagram_path, dispense_diagram_path = create_safety_diagrams(
        PLANNING_DIAGRAM_PATH,
        DISPENSE_DIAGRAM_PATH,
    )
    doc = Document()
    configure_styles(doc)
    for section in doc.sections:
        configure_page(section)
        configure_header_footer(section)
    bullet_id, protocol_decimal_id, dispense_decimal_id = add_custom_numbering(doc)

    # First-page masthead
    p = doc.add_paragraph(style="Kicker")
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(7)
    p.add_run("SAFETY CONTROL SPECIFICATION  /  VERSION 1.0")

    p = doc.add_paragraph(style="Document Title")
    p.add_run("Agent Safety Architecture")
    p = doc.add_paragraph(style="Document Subtitle")
    p.add_run("Guardrails and Gatekeepers for Personalized Supplement Protocols, Cartridge Ordering, and Pill Dispensing")

    metadata = [
        ("Purpose", "Define deterministic safety controls around agent recommendations and physical actions"),
        ("Scope", "Protocol approval, cartridge order preparation, per-pill dispensing, and cartridge validation"),
        ("Safety posture", "Fail closed; stale, missing, conflicting, or unverifiable data blocks action"),
        ("Freshness baseline", "Profile evidence must be confirmed within the preceding 90 days"),
        ("Status", "Draft system requirements for product, clinical, engineering, and quality review"),
    ]
    for label, value in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2.5)
        p.paragraph_format.line_spacing = 1.0
        r1 = p.add_run(f"{label}: ")
        set_run_font(r1, size=10.5, color=NAVY, bold=True)
        r2 = p.add_run(value)
        set_run_font(r2, size=10.5, color=BLACK)

    add_callout(
        doc,
        "Core rule",
        "The agent may propose a protocol, but it may not approve, order, or dispense anything unless the relevant gate produces an explicit ALLOW decision from current, complete, and trusted evidence.",
    )

    heading(doc, "1. Executive Summary", 1)
    add_para(
        doc,
        "The system uses defense in depth. Guardrails constrain what the agent can recommend and how it reasons about user data. Gatekeepers sit at irreversible action boundaries - protocol approval, cartridge-order creation, and pill dispensing - and independently decide whether the requested action is permitted. The language model may explain or propose; deterministic services own the final safety decision.",
    )
    add_para(
        doc,
        "Three control mechanisms are required. First, every proposed multi-supplement protocol is evaluated against medical contraindications, medication interactions, body-weight-based dose limits, and schedule conflicts. Second, the same safety logic is re-evaluated for the individual pill immediately before release. Third, the hardware or dispensing service validates that the installed cartridge belongs to the intended user and blocks cartridges associated with a quarantined or recalled lot, expired pills, or an invalid order manifest.",
    )

    heading(doc, "2. Safety Principles", 1)
    principles = [
        "Fail closed. Unknown, missing, stale, malformed, or contradictory evidence is a blocking condition, not an invitation to infer.",
        "Fresh evidence. Relevant profile statements, medication status, and body weight must have a trustworthy confirmation timestamp no older than 90 days at the moment of evaluation.",
        "Deterministic enforcement. A versioned rules engine and trusted data services calculate the decision. The agent cannot override, reinterpret, or suppress a failed check.",
        "Defense in depth. Approval at planning time does not replace order-time or dispense-time validation.",
        "User-specific binding. Every approval, order, cartridge, and dispense event is cryptographically or logically bound to one user, one protocol version, and one ruleset version.",
        "Traceability. Every decision stores inputs, timestamps, rule identifiers, outcomes, and the reason shown to the user or operator.",
        "Clinical governance. Medical rules and dose formulas are maintained in a clinician-approved, version-controlled knowledge base. This document specifies enforcement behavior, not clinical content.",
    ]
    for item in principles:
        add_list_item(doc, item, bullet_id)

    heading(doc, "3. Control Model and Decision Ownership", 1)
    add_para(
        doc,
        "Guardrails and gatekeepers serve different purposes. Guardrails shape proposals and conversation. Gatekeepers authorize actions. A gatekeeper must receive structured inputs and return a structured decision; free-text agent output is never treated as approval.",
    )
    figure = doc.add_paragraph()
    figure.alignment = WD_ALIGN_PARAGRAPH.CENTER
    figure.paragraph_format.space_before = Pt(4)
    figure.paragraph_format.space_after = Pt(3)
    figure.paragraph_format.keep_with_next = True
    image_run = figure.add_run()
    inline_shape = image_run.add_picture(str(planning_diagram_path), width=Inches(6.35))
    inline_shape._inline.docPr.set("name", "Protocol and order authorization flow")
    inline_shape._inline.docPr.set(
        "descr",
        "Flow diagram showing an agent proposal passing through a conversational guardrail and deterministic protocol safety gate. Passing checks continue to the order gatekeeper; failed checks route to input collection, replanning, clinical review, or rejection.",
    )
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(0)
    caption.paragraph_format.space_after = Pt(10)
    caption.paragraph_format.keep_with_next = False
    caption_run = caption.add_run("Figure 1. Protocol and cartridge-order authorization flow.")
    set_run_font(caption_run, size=9, color=MUTED, italic=True)

    figure = doc.add_paragraph()
    figure.alignment = WD_ALIGN_PARAGRAPH.CENTER
    figure.paragraph_format.space_before = Pt(2)
    figure.paragraph_format.space_after = Pt(3)
    figure.paragraph_format.keep_with_next = True
    image_run = figure.add_run()
    inline_shape = image_run.add_picture(str(dispense_diagram_path), width=Inches(6.35))
    inline_shape._inline.docPr.set("name", "Dispense-time defense flow")
    inline_shape._inline.docPr.set(
        "descr",
        "Flow diagram showing authenticated user and installed cartridge inputs splitting into cartridge and per-pill gatekeepers. Both gates must allow before one pill is released; any failed gate produces a hard block. The cartridge gate includes a cloud DHR SAFE-status check.",
    )
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(0)
    caption.paragraph_format.space_after = Pt(10)
    caption.paragraph_format.keep_with_next = True
    caption_run = caption.add_run("Figure 2. Both independent dispense-time gatekeepers must allow physical release.")
    set_run_font(caption_run, size=9, color=MUTED, italic=True)
    add_table(
        doc,
        ["Control layer", "Responsibility", "Permitted output", "Cannot do"],
        [
            ["Agent guardrail", "Constrain proposal content; collect missing data; explain decisions", "Draft protocol, question, or explanation", "Approve its own recommendation or bypass a failed rule"],
            ["Protocol gatekeeper", "Validate the complete proposed protocol", "ALLOW, PENDING_INPUT, REPLAN, CLINICAL_REVIEW, or REJECT", "Create an order before ALLOW"],
            ["Order gatekeeper", "Bind approved protocol to a cartridge order manifest", "Signed order authorization or block", "Order an unapproved or changed protocol"],
            ["Dispense gatekeeper", "Re-check one pill against current user and schedule state", "One-time dispense authorization or block", "Rely solely on the original protocol approval"],
            ["Cartridge gatekeeper", "Validate user, cartridge, lot, expiry, and manifest", "Cartridge enablement or hard block", "Enable an unbound, recalled, quarantined, or expired cartridge"],
        ],
        [1500, 2920, 2500, 2440],
        font_size=9,
    )

    heading(doc, "3.1 Canonical Decision States", 2)
    add_table(
        doc,
        ["State", "Meaning", "Required next action"],
        [
            ["ALLOW", "All required checks passed using current trusted evidence", "Proceed only for the exact action and version authorized"],
            ["PENDING_USER_INPUT", "A required answer or profile datum is missing or stale", "Ask a targeted question; store the timestamped answer; re-run all checks"],
            ["REPLAN", "The contents may be acceptable but dose or timing is unsafe or inconsistent", "Generate a corrected proposal; do not silently change an approved protocol"],
            ["CLINICAL_REVIEW", "A condition, medication, dose, ambiguity, or policy requires qualified review", "Escalate; prevent automated approval and dispensing"],
            ["REJECT", "The requested action is prohibited, invalid, mismatched, recalled, or expired", "Block the action and present a safe, non-sensitive reason"],
        ],
        [1800, 3720, 3840],
        font_size=9.2,
    )

    heading(doc, "3.2 Freshness and Evidence Rules", 2)
    add_para(
        doc,
        "For this specification, 'three months' is implemented as a configurable 90-day freshness window. The gate compares the evidence confirmation timestamp to a trusted server timestamp. A record exactly 90 days old remains valid; a record older than 90 days is stale. Product and clinical governance may replace this with a calendar-month definition, but one definition must be used consistently across every gate.",
    )
    add_list_item(doc, "A negative answer is valid only when it names the relevant condition or medication class, identifies the user, records the source, and carries a confirmation timestamp.", bullet_id)
    add_list_item(doc, "Absence of a record is not equivalent to a negative answer. Imported or inferred data cannot silently substitute for explicit user confirmation when the rule requires it.", bullet_id)
    add_list_item(doc, "Any material profile change invalidates prior approvals that depended on the changed field, even if the previous answer is still within 90 days.", bullet_id)

    heading(doc, "4. Mechanism One - Protocol Safety Gate", 1)
    add_para(
        doc,
        "The agent proposes a protocol containing one or more supplements. Each protocol item must include an ingredient or product identifier, strength, units per dose, calculated daily exposure, proposed administration times, duration, and rule-source version. The gate evaluates every item individually and the protocol as a whole.",
    )

    heading(doc, "4.1 Required Evaluation Sequence", 2)
    steps = [
        "Normalize the proposal into structured protocol items and reject unknown ingredients, ambiguous units, unsupported formulations, or incomplete schedules.",
        "Load the user's current profile, including medical-condition answers, medication list or explicit non-use confirmations, body weight, timezone, existing supplement schedule, and timestamps.",
        "Load the active clinician-approved ruleset for contraindications, restrictions, interactions, dose formulas, upper limits, timing windows, and escalation criteria.",
        "Evaluate condition restrictions, medication interactions, dose appropriateness, and schedule conflicts for every protocol item.",
        "Evaluate whole-protocol cumulative exposure, duplicate ingredients, and cross-item timing constraints.",
        "Return one canonical decision with machine-readable reason codes. ALLOW is possible only if every mandatory check passes.",
    ]
    for step in steps:
        add_list_item(doc, step, protocol_decimal_id)

    heading(doc, "4.2 Existing Medical Conditions", 2)
    add_para(
        doc,
        "When the ruleset marks an ingredient as prohibited or restricted for a condition - for example, hypertension - the system must locate a recent, explicit user response for that condition. A negative answer such as 'No, I do not have blood-pressure problems' is sufficient only if it is tied to the relevant condition and was confirmed within the preceding 90 days.",
    )
    add_table(
        doc,
        ["Profile evidence", "Gate result", "System behavior"],
        [
            ["Recent explicit negative answer", "Continue", "Evaluate the remaining rules"],
            ["No answer or answer older than 90 days", "PENDING_USER_INPUT", "Ask the condition-specific question; do not approve until answered"],
            ["Positive answer", "CLINICAL_REVIEW or REJECT", "Apply the clinician-authored rule; the agent cannot override it"],
            ["Contradictory or uncertain answer", "CLINICAL_REVIEW", "Resolve the discrepancy through the approved workflow"],
        ],
        [2700, 1900, 4760],
        font_size=9.3,
    )
    add_callout(
        doc,
        "Privacy-aware questioning",
        "Ask only the minimum question required by the triggered rule. The user-facing message should explain why the answer is needed without revealing unnecessary internal rule details.",
        fill=LIGHT_GRAY,
        accent=MUTED,
    )

    heading(doc, "4.3 Medication Interactions", 2)
    add_para(
        doc,
        "The gate screens the proposed supplement against prescription medicines, over-the-counter medicines, and any medication classes named by the active rule. The user's medication status must have been confirmed within 90 days. A stale list, an incomplete reconciliation, or uncertainty about use blocks approval.",
    )
    for item in [
        "If the profile confirms that the user does not take the interacting medication or class, and the confirmation is current, continue.",
        "If medication status is missing or stale, return PENDING_USER_INPUT and request an updated medication reconciliation before re-evaluating the entire protocol.",
        "If an interacting medication is present, apply the configured action: REPLAN, CLINICAL_REVIEW, or REJECT. Do not rely on the agent to judge interaction severity.",
        "If the user changes any medication after approval, invalidate the approval token and require a fresh protocol evaluation.",
    ]:
        add_list_item(doc, item, bullet_id)

    heading(doc, "4.4 Body-Weight-Based Dose Validation", 2)
    add_para(
        doc,
        "When a recommended daily dose depends on body weight, the gate must retrieve a weight measurement confirmed within 90 days and calculate the permitted range using the active rule's units, age restrictions, rounding behavior, minimum and maximum values, and absolute upper limit. The calculation must consider total daily exposure from all protocol items and existing supplements that contain the same active ingredient.",
    )
    add_table(
        doc,
        ["Check", "Requirement", "Failure outcome"],
        [
            ["Weight freshness", "Measurement timestamp is within 90 days", "PENDING_USER_INPUT"],
            ["Unit integrity", "Weight and formula units are explicit and safely converted", "REJECT if ambiguous; otherwise recalculate"],
            ["Formula validity", "Clinician-approved formula applies to the user's supported population", "CLINICAL_REVIEW"],
            ["Daily exposure", "Proposed amount plus all known sources remains within the permitted range and upper limit", "REPLAN or REJECT"],
            ["Rounding", "Rounding follows the rule and available pill strengths without exceeding limits", "REPLAN"],
        ],
        [1900, 4840, 2620],
        font_size=9.3,
    )
    add_callout(
        doc,
        "No silent assumptions",
        "The system must not substitute an old weight, a population average, an estimated weight, or a value from another user. If the valid calculation cannot be completed, approval is blocked.",
        fill=PALE_GOLD,
        accent=GOLD,
        color=GOLD,
    )

    heading(doc, "4.5 Administration-Time Conflicts", 2)
    add_para(
        doc,
        "The schedule engine evaluates the proposed times against the user's existing supplement schedule and all timing rules. A timing rule is represented as a directed separation constraint, a prohibited window, or a required context such as with food. For example, if a rule requires a one-hour window after iron, the engine must not schedule the conflicting supplement until at least 60 minutes after the recorded or planned iron dose.",
    )
    for item in [
        "Evaluate the complete local-day schedule in the user's timezone, including windows that cross midnight.",
        "Treat equal-to-boundary timing according to the rule definition; for an 'at least 60 minutes' rule, exactly 60 minutes is valid.",
        "If a safe schedule exists, return REPLAN with the proposed corrected times for user confirmation.",
        "If no compliant schedule exists, return CLINICAL_REVIEW or REJECT according to policy.",
        "Store timing rules by identifiers and durations; do not encode medical timing knowledge only in prompts or free text.",
    ]:
        add_list_item(doc, item, bullet_id)

    heading(doc, "4.6 Protocol Approval Artifact", 2)
    add_para(
        doc,
        "An ALLOW decision produces a signed, immutable approval artifact containing the user ID, protocol ID and version, normalized items and schedule, profile snapshot identifiers, ruleset version, evaluation timestamp, freshness expiry, and decision hash. Any change to the protocol, relevant profile data, ruleset, recall status, or user identity invalidates the artifact and forces re-evaluation.",
    )

    heading(doc, "5. Cartridge Order Preparation Gate", 1)
    add_para(
        doc,
        "After protocol approval, the agent may prepare an order request; the order service creates the physical order only after validating the approval artifact. The agent supplies intent and explanatory context, while the order gatekeeper owns authorization.",
    )
    heading(doc, "5.1 Order Preconditions", 2)
    for item in [
        "The approval artifact is authentic, unexpired, and bound to the same user and protocol version.",
        "No relevant profile field, medication status, body weight, ruleset, product status, or recall status has changed since approval.",
        "Every ordered product and strength matches an approved protocol item; quantities and duration are within the authorized scope.",
        "The generated cartridge manifest identifies user, order, protocol, product, strength, pill count, lot eligibility requirements, intended administration schedule, and expiration constraints.",
        "A second deterministic check confirms that the order does not introduce substitutions, duplicates, or packaging changes that alter dose or timing behavior.",
    ]:
        add_list_item(doc, item, bullet_id)
    add_callout(
        doc,
        "Atomicity requirement",
        "The order is created only if all cartridges and pills in the order are authorized. Partial creation is permitted only when an explicit product policy defines a safe, user-visible partial-order workflow.",
        fill=LIGHT_GRAY,
        accent=MUTED,
    )

    heading(doc, "5.2 Order Output and Audit", 2)
    add_para(
        doc,
        "The output is a signed order authorization and a versioned cartridge manifest. The audit event records the originating agent request, approval hash, rule version, product identifiers, quantities, intended user, timestamp, result, and reason codes. A human-readable explanation is stored separately from the safety decision so wording changes cannot alter authorization semantics.",
    )

    heading(doc, "6. Mechanism Two - Per-Pill Dispense Gate", 1)
    add_para(
        doc,
        "Immediately before a pill is physically released, the system repeats the applicable protocol checks for that individual pill using the newest available state. This is not a lightweight confirmation of the previous decision; it is a new authorization for one specific pill, user, cartridge, and time window.",
    )
    heading(doc, "6.1 Dispense-Time Inputs", 2)
    add_table(
        doc,
        ["Input", "Required validation"],
        [
            ["User identity", "Authenticated active user matches protocol, order, device session, and cartridge binding"],
            ["Pill identity", "Product, ingredient, strength, lot, and expiry match the manifest and approved protocol"],
            ["Current profile", "Relevant conditions, medication status, and body weight remain complete and within 90 days"],
            ["Current schedule state", "Actual and planned doses, daily cumulative exposure, previous-dispense events, and timezone are current"],
            ["Safety state", "Ruleset, recalls, lot quarantine, product status, and device integrity checks are current"],
        ],
        [2200, 7160],
        font_size=9.4,
    )

    heading(doc, "6.2 Per-Pill Decision Logic", 2)
    steps = [
        "Validate user-to-device and user-to-cartridge identity binding.",
        "Validate the pill against the signed order manifest, lot status, and expiration date.",
        "Re-run relevant condition and medication interaction rules using current profile evidence.",
        "Recalculate the user's total daily exposure including the pill about to be dispensed; block duplicate, excess, or too-soon dosing.",
        "Evaluate actual timing, not merely the planned schedule. For the iron example, measure the required interval from the recorded iron dose before authorizing the next pill.",
        "Issue a single-use, short-lived dispense authorization bound to the exact actuator command. Consume the authorization atomically when the pill is released.",
    ]
    for step in steps:
        add_list_item(doc, step, dispense_decimal_id)

    add_callout(
        doc,
        "Hard safety boundary",
        "If the safety service, trusted clock, identity service, lot-status service, or required profile data is unavailable, the device must not dispense. A user-safe retry or escalation path may be offered, but availability does not outrank safety.",
        fill=PALE_RED,
        accent=RED,
        color=RED,
    )

    heading(doc, "6.3 Concurrency and Duplicate Prevention", 2)
    add_para(
        doc,
        "The system must prevent two devices, retries, or concurrent requests from authorizing the same dose. A dispense ledger and idempotency key reserve the dose before actuation and finalize it only after confirmed release. Ambiguous actuator outcomes are treated as safety events requiring reconciliation; the system must not automatically dispense a replacement when it cannot determine whether a pill was released.",
    )

    heading(doc, "6.4 User Communication", 2)
    add_para(
        doc,
        "A block message should state what happened, what the user can do next, and whether professional review is required. It must not imply a diagnosis or expose sensitive medication or condition data on a shared device. Examples include: 'Your profile information needs to be updated before this dose can be released,' 'This dose is too close to another scheduled supplement,' and 'This cartridge cannot be used; please contact support.'",
    )

    heading(doc, "7. Mechanism Three - Cartridge Safety Gate", 1)
    add_para(
        doc,
        "A cartridge is enabled only when its digital identity, physical contents, and lifecycle status match the intended user and order. These checks occur at installation, at session start, and again immediately before dispensing when status may have changed.",
    )

    heading(doc, "7.1 User-to-Cartridge Matching", 2)
    for item in [
        "Read a tamper-resistant cartridge identifier, such as a signed barcode, NFC tag, or secure serial number.",
        "Query the associated DHR in the cloud and verify that its current status is SAFE before continuing. A non-SAFE, unknown, stale, or unavailable status is a hard blocking condition.",
        "Resolve the identifier to the signed manifest and verify its signature, issuer, order ID, intended user ID, product content, and lifecycle status.",
        "Compare the intended user with the authenticated active user. A mismatch is a hard REJECT; no confirmation dialog may override it.",
        "Prevent reassignment or reuse unless a controlled operational workflow explicitly reissues and re-signs the cartridge binding.",
        "Record insertion, removal, mismatch, enablement, and tamper events in the audit trail.",
    ]:
        add_list_item(doc, item, bullet_id)

    heading(doc, "7.2 Lot Recall or Quarantine", 2)
    add_para(
        doc,
        "Before enabling or dispensing from a cartridge, the system checks the current status of every contained lot. Status values should include RELEASED, QUARANTINED, RECALLED, and UNKNOWN. Only RELEASED is eligible. QUARANTINED, RECALLED, and UNKNOWN produce a hard block. The system must be able to distribute urgent lot-status updates and invalidate previously cached permissions.",
    )

    heading(doc, "7.3 Expiration and Shelf-Life", 2)
    add_para(
        doc,
        "The system validates the manufacturer's expiration date and any applicable cartridge beyond-use date against a trusted local date and timezone policy. Pills must not be dispensed after expiration. If expiry occurs during a multi-dose period, only doses strictly before the configured cutoff remain eligible. Unreadable, missing, contradictory, or untrusted dates are treated as invalid.",
    )

    heading(doc, "7.4 Cartridge Decision Matrix", 2)
    add_table(
        doc,
        ["Check", "Pass condition", "Failure behavior"],
        [
            ["User binding", "Active user equals manifest user", "Hard block; log mismatch"],
            ["Manifest integrity", "Signature and contents validate", "Hard block; quarantine cartridge"],
            ["Lot status", "Every lot is RELEASED", "Hard block for QUARANTINED, RECALLED, or UNKNOWN"],
            ["Expiry", "Current time is before all applicable cutoffs", "Hard block; mark cartridge unusable"],
            ["Content match", "Product, strength, count, and lot match manifest", "Hard block; operational investigation"],
            ["Tamper state", "Seal and secure identity checks pass", "Hard block; quarantine and log"],
        ],
        [1900, 3480, 3980],
        font_size=9.2,
    )

    heading(doc, "8. Data, Audit, and Operational Requirements", 1)
    heading(doc, "8.1 Minimum Structured Data", 2)
    add_table(
        doc,
        ["Entity", "Minimum fields"],
        [
            ["Profile evidence", "User ID, question or field ID, normalized value, source, confirmation timestamp, provenance, supersession status"],
            ["Protocol item", "Ingredient/product ID, formulation, strength, units per dose, daily exposure, schedule, duration, rule references"],
            ["Safety decision", "Gate, action, decision state, reason codes, input snapshot hashes, ruleset version, trusted timestamp, expiry"],
            ["Order manifest", "Order and user IDs, approval hash, protocol version, products, strengths, counts, lot constraints, schedule, signature"],
            ["Cartridge record", "Secure identifier, order/user binding, products, counts, lot IDs, expiration and beyond-use dates, status"],
            ["Dispense event", "Authorization ID, user/device/cartridge/pill IDs, planned and actual time, pre-check result, actuator result, idempotency key"],
        ],
        [2100, 7260],
        font_size=9.1,
    )

    heading(doc, "8.2 Auditability", 2)
    for item in [
        "Logs are append-only, access-controlled, and tamper-evident, with retention defined by applicable policy and regulation.",
        "Every block and approval includes stable reason codes suitable for monitoring, testing, investigation, and user-support workflows.",
        "Sensitive health data is minimized in logs and user-facing device messages. Detailed reasons are available only to authorized roles.",
        "Rule and model versions are recorded independently. Changing the model prompt cannot alter deterministic gate behavior.",
        "Clock drift, offline operation, stale caches, rule-service failure, and lot-status uncertainty are observable safety events.",
    ]:
        add_list_item(doc, item, bullet_id)

    heading(doc, "8.3 Fail-Safe Operational Behavior", 2)
    add_para(
        doc,
        "The system must define recovery behavior without weakening the gate. Safe responses include refreshing data, asking the user to reconfirm a profile field, re-planning a schedule, routing to clinical review, quarantining a cartridge, or contacting support. Unsafe responses include guessing, treating silence as a negative answer, using a stale cached ALLOW decision, changing the user's dose without confirmation, or allowing a device operator to bypass a hard block.",
    )

    heading(doc, "9. Reference Decision Algorithm", 1)
    add_para(doc, "The following pseudocode captures the required enforcement order. It is illustrative; production logic should use typed inputs, policy identifiers, and atomic transactions.")
    code_lines = [
        "evaluate(action, user, protocol, pill, cartridge, now):",
        "  require trusted_identity(user) and trusted_clock(now)",
        "  require current_ruleset() and complete_structured_inputs(action)",
        "  require relevant_profile_evidence_is_current(user, max_age=90_days)",
        "  require condition_rules_pass(user, protocol or pill)",
        "  require medication_interaction_rules_pass(user, protocol or pill)",
        "  require weight_based_total_daily_dose_passes(user, protocol or pill)",
        "  require actual_and_planned_timing_rules_pass(user, protocol or pill, now)",
        "  if action in {ORDER, DISPENSE}:",
        "    require valid_approval_binding(user, protocol)",
        "  if action == DISPENSE:",
        "    require cartridge_user_lot_expiry_manifest_checks_pass(cartridge, user, pill, now)",
        "    require reserve_single_use_authorization_atomically()",
        "  return ALLOW only if every required predicate is true; otherwise block with reason codes",
    ]
    for line in code_lines:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.22)
        p.paragraph_format.right_indent = Inches(0.12)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p_pr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "F7F8FA")
        p_pr.append(shd)
        r = p.add_run(line)
        set_run_font(r, name="Courier New", size=8.9, color=NAVY)

    heading(doc, "10. Minimum Acceptance Tests", 1)
    add_para(doc, "The implementation is not ready for release until at least the following end-to-end cases pass in automated tests and controlled hardware tests.")
    add_table(
        doc,
        ["ID", "Scenario", "Expected result"],
        [
            ["AT-01", "Condition rule is triggered; explicit negative answer is 30 days old", "Condition check passes; remaining checks still run"],
            ["AT-02", "Condition answer is missing or 91 days old", "PENDING_USER_INPUT; no approval or order"],
            ["AT-03", "User reports the restricted condition", "CLINICAL_REVIEW or REJECT according to rule; no agent override"],
            ["AT-04", "Medication list is current and includes an interacting medicine", "Configured interaction action; no approval"],
            ["AT-05", "Weight-based dose uses weight older than 90 days", "PENDING_USER_INPUT; dose is not estimated"],
            ["AT-06", "Combined ingredients exceed daily upper limit", "REPLAN or REJECT; cumulative exposure shown in audit"],
            ["AT-07", "Conflicting supplement is scheduled 45 minutes after iron; rule requires 60", "REPLAN or block; no approval at 45 minutes"],
            ["AT-08", "Conflicting supplement is actually requested exactly 60 minutes after iron", "Timing check passes if all other checks pass"],
            ["AT-09", "Approved protocol is changed before order creation", "Approval artifact invalid; order blocked"],
            ["AT-10", "Medication profile changes after protocol approval", "Approval invalidated; order and dispense blocked pending re-evaluation"],
            ["AT-11", "Correct cartridge is installed for a different authenticated user", "Hard block; mismatch event logged"],
            ["AT-12", "Cartridge lot changes from RELEASED to RECALLED after installation", "Next dispense blocked immediately; cartridge disabled"],
            ["AT-13", "Pill is expired or expiry is unreadable", "Hard block; cartridge marked unusable or quarantined"],
            ["AT-14", "Two devices request the same dose concurrently", "Only one single-use authorization succeeds"],
            ["AT-15", "Safety or lot-status service is unavailable", "No dispense; safe retry or escalation message"],
        ],
        [900, 5180, 3280],
        font_size=8.7,
    )

    heading(doc, "11. Governance and Open Decisions", 1)
    add_para(
        doc,
        "Before implementation, product, clinical, quality, privacy, security, and engineering owners should approve the rule-source governance model, exact definition of the three-month window, escalation pathways, offline behavior, user-authentication strength, lot-status distribution latency, cartridge identity technology, expiration cutoff semantics, and jurisdiction-specific compliance requirements.",
    )
    add_callout(
        doc,
        "Release criterion",
        "No generative component may be placed on the authorization path. Safety decisions must remain reproducible from stored structured inputs and the recorded ruleset version.",
        fill=PALE_GREEN,
        accent=GREEN,
        color=GREEN,
    )

    # Core properties and compatibility.
    doc.core_properties.title = "Agent Safety Architecture: Guardrails and Gatekeepers"
    doc.core_properties.subject = "Safety controls for supplement protocols, cartridge ordering, and pill dispensing"
    doc.core_properties.author = "BQ Product and Safety"
    doc.core_properties.keywords = "agent safety, guardrails, gatekeepers, supplements, protocol, dispensing, cartridge"
    doc.core_properties.comments = "Draft for multidisciplinary review"
    doc.settings.element.append(OxmlElement("w:updateFields"))
    doc.settings.element[-1].set(qn("w:val"), "true")

    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_document())
